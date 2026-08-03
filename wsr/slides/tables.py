"""
Data table slides: pending DCRs (5–6+), DDP, handoff.

Pending eval/impl tables are split across multiple slides when needed so the
table never overflows a single page.
"""

from __future__ import annotations

from pptx import Presentation

from wsr.constants import PENDING_ROWS_PER_SLIDE
from wsr.slides.base import add_table, empty_row, new_content_slide


def _chunked(items: list, size: int) -> list[list]:
    if not items:
        return [[]]
    return [items[i : i + size] for i in range(0, len(items), size)]


def add_pending_slide(
    prs: Presentation,
    title: str,
    report_date: str,
    slide_number: int,
    items: list[dict],
    mode: str,
    *,
    rows_per_slide: int = PENDING_ROWS_PER_SLIDE,
) -> int:
    """
    Add one or more pending-table slides.

    Returns how many slides were added so the deck can renumber later slides.
    """
    if mode == "evaluation":
        headers = [
            "Sr No",
            "DCR ID",
            "Summary",
            "Current Status",
            "Eval Closure date",
            "Support Required",
        ]
    else:
        headers = [
            "Sr No",
            "DCR ID",
            "Summary",
            "Current Status",
            "Impl Closure Date",
            "Support Required",
        ]
    widths = [0.58, 0.9, 3.4, 1.85, 1.45, 3.0]

    chunks = _chunked(items, max(1, rows_per_slide))
    total_pages = len(chunks)
    row_offset = 0

    for page_index, chunk in enumerate(chunks):
        page_title = title if total_pages == 1 else f"{title} ({page_index + 1}/{total_pages})"
        slide = new_content_slide(prs, page_title, report_date, slide_number + page_index)

        rows = [
            [
                str(row_offset + i + 1),
                str(item["dcr_id"]),
                item["summary"],
                item["status"],
                item["closure_date"],
                item["support"],
            ]
            for i, item in enumerate(chunk)
        ]
        if not rows:
            rows = [empty_row(len(headers))]

        add_table(slide, headers, rows, col_widths=widths)
        row_offset += len(chunk)

    return total_pages


def add_ddp_slide(
    prs: Presentation,
    report_date: str,
    items: list[dict] | None = None,
    *,
    slide_number: int = 7,
) -> None:
    """DDP MS 4-5 rows from Non STLA (Is DDP Testing Needed = Yes)."""
    slide = new_content_slide(prs, "PFS DDP Details MS 4-5", report_date, slide_number)
    headers = [
        "Sr.No",
        "DCR ID",
        "DCR Summary",
        "Plan Dates",
        "Appeared Dates",
        "Program",
        "Dependencies",
        "Remarks",
    ]
    rows = [
        [
            str(i + 1),
            item.get("dcr_id", ""),
            item.get("summary", ""),
            item.get("plan_date", ""),
            item.get("appeared_date", ""),
            item.get("program", ""),
            item.get("dependencies", ""),
            item.get("remarks", ""),
        ]
        for i, item in enumerate(items or [])
    ]
    if not rows:
        rows = [empty_row(len(headers))]
    add_table(
        slide,
        headers,
        rows,
        col_widths=[0.55, 0.9, 3.2, 1.2, 1.25, 0.9, 2.2, 2.5],
    )


def add_handoff_slide(
    prs: Presentation,
    report_date: str,
    items: list[dict] | None = None,
    *,
    slide_number: int = 8,
) -> None:
    del items
    slide = new_content_slide(prs, "Q3-2026 – Eval Handoff from onsite", report_date, slide_number)
    headers = ["Sr. No.", "DCR ID", "Summary", "Evaluator", "Eval Handoff Date", "Remark"]
    rows = [[""] * len(headers)]
    add_table(slide, headers, rows, col_widths=[0.62, 0.9, 3.6, 1.05, 1.55, 2.6])
